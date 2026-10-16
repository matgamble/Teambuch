import sys
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY=RGBColor(0x12,0x26,0x4a); GOLD=RGBColor(0x8a,0x6d,0x2b)
d=Document(); s=d.sections[0]
s.page_width,s.page_height=Mm(210),Mm(297)
s.left_margin=s.right_margin=Mm(22); s.top_margin=Mm(14); s.bottom_margin=Mm(14)
st=d.styles['Normal']; st.font.name='Georgia'; st.font.size=Pt(11); st.font.color.rgb=NAVY
st.element.rPr.rFonts.set(qn('w:eastAsia'),'Georgia')
st.paragraph_format.space_after=Pt(7); st.paragraph_format.line_spacing=1.15

def border(el_pr, tag, sz, color):
    b=el_pr.find(qn('w:pBdr')) if tag=='p' else None
def set_cell_borders(cell, sz=16, color='12264A', fill=None):
    tcPr=cell._tc.get_or_add_tcPr(); b=OxmlElement('w:tcBorders')
    for e in ('top','left','bottom','right'):
        x=OxmlElement(f'w:{e}'); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),str(sz)); x.set(qn('w:color'),color); b.append(x)
    tcPr.append(b)
    if fill:
        sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),fill); tcPr.append(sh)
def bottom_line(p,color='C9A85C',sz=12):
    pPr=p._p.get_or_add_pPr(); b=OxmlElement('w:pBdr'); x=OxmlElement('w:bottom')
    x.set(qn('w:val'),'single'); x.set(qn('w:sz'),str(sz)); x.set(qn('w:space'),'4'); x.set(qn('w:color'),color); b.append(x); pPr.append(b)
def top_line(p):
    pPr=p._p.get_or_add_pPr(); b=OxmlElement('w:pBdr'); x=OxmlElement('w:top')
    x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'6'); x.set(qn('w:space'),'4'); x.set(qn('w:color'),'C9A85C'); b.append(x); pPr.append(b)
def no_borders(t):
    tblPr=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        x=OxmlElement(f'w:{e}'); x.set(qn('w:val'),'nil'); b.append(x)
    tblPr.append(b)

# Kopf: Logo links, Datum rechts
t=d.add_table(rows=1,cols=2); no_borders(t); t.autofit=False
w=[Mm(110),Mm(56)]
for c,x in zip(t.rows[0].cells,w): c.width=x
t.rows[0].cells[0].paragraphs[0].add_run().add_picture('logo.png',height=Mm(24))
p=t.rows[0].cells[1].paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=p.add_run('Augsburg, im Oktober 2026'); r.font.size=Pt(10)
bottom_line(d.add_paragraph())

p=d.add_paragraph(); p.paragraph_format.space_before=Pt(14); p.paragraph_format.space_after=Pt(2)
r=p.add_run('Herzlichen Glückwunsch zur Wahl in den Elternbeirat!'); r.bold=True; r.font.size=Pt(19)
p=d.add_paragraph(); r=p.add_run('Elternbeirat der Kita Reischlestraße · Kita-Jahr 2026/27'); r.italic=True; r.font.color.rgb=GOLD
p.paragraph_format.space_after=Pt(12)

d.add_paragraph(f'Liebe {sys.argv[1]},')
d.add_paragraph('Du wurdest für das Kita-Jahr 2026/27 in den Elternbeirat der Kita Reischlestraße gewählt – dazu gratulieren wir dir ganz herzlich!')
p=d.add_paragraph(); p.add_run('Wir freuen uns sehr, dass du diese Aufgabe übernimmst,')
p.add_run(' und auf eine vertrauensvolle, offene Zusammenarbeit. Der Elternbeirat ist die Stimme der Eltern in unserer Kita, ein wichtiges Bindeglied zwischen Familien, Team und Leitung und damit eine große Bereicherung für unser Haus.')
d.add_paragraph('Vielen Dank für deine Zeit und dein Engagement. Gemeinsam gestalten wir das Kita-Jahr für unsere Kinder.')

sp=d.add_paragraph(); sp.paragraph_format.space_after=Pt(2)
t=d.add_table(rows=1,cols=2); t.autofit=False; t.alignment=1
lay=OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); t._tbl.tblPr.append(lay)
t.columns[0].width=Mm(50); t.columns[1].width=Mm(116)
a,b=t.rows[0].cells; a.width=Mm(50); b.width=Mm(116)
for c in (a,b): set_cell_borders(c,fill='FAF6EA')
a.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
a.paragraphs[0].add_run().add_picture('qr-whatsapp-elternbeirat.png',width=Mm(42))
for c in (a,b):
    tcPr=c._tc.get_or_add_tcPr(); v=OxmlElement('w:vAlign'); v.set(qn('w:val'),'center'); tcPr.append(v)
p=b.paragraphs[0]; r=p.add_run('Unsere WhatsApp-Gruppe für den Elternbeirat'); r.bold=True; r.font.size=Pt(12)
p=b.add_paragraph(); r=p.add_run('Zur einfachen Abstimmung von Terminen und für den kurzen Draht untereinander gibt es die Gruppe '); r.font.size=Pt(10)
r=p.add_run('„Elternbeirat KiTa Reischle 26/27“'); r.bold=True; r.font.size=Pt(10); p.add_run('.').font.size=Pt(10)
p=b.add_paragraph(); r=p.add_run('So trittst du bei: WhatsApp öffnen, mit der Kamera den QR-Code scannen und dem Beitritt zustimmen. Bitte gib den Code nicht weiter, damit die Gruppe im Kreis des Elternbeirats bleibt.'); r.font.size=Pt(10)

p=d.add_paragraph(); p.paragraph_format.space_before=Pt(16)
p.add_run('Mit herzlichen Grüßen')
for _ in range(2): d.add_paragraph().paragraph_format.space_after=Pt(0)
p=d.add_paragraph('Matthias Stilwell'); p.paragraph_format.space_after=Pt(0)
d.add_paragraph('Kita-Leitung')

f=s.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; top_line(f)
r=f.add_run('Kita Reischlestraße · Gemeinsam. Für Kinder. Mit Herz und Ideen.'); r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(0x5b,0x6b,0x85)
d.save(f'einzelbriefe/elternbeirat-brief-{sys.argv[1].lower()}.docx')
